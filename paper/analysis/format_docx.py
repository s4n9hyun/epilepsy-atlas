"""Epilepsia review formatting for the pandoc Word output: Times New Roman 12 pt, double spacing,
continuous line numbers and page numbers."""
import sys

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = 1
    run = p.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def main(path, line_numbers=True):
    doc = Document(path)
    for st in doc.styles:
        if st.type == 1:  # paragraph styles
            st.font.name = "Times New Roman"
            st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman") if st.element.rPr is not None and st.element.rPr.rFonts is not None else None
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    for p in doc.paragraphs:
        if p.style.name.startswith(("Heading", "Title")):
            continue
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    for section in doc.sections:
        if line_numbers:
            ln = OxmlElement("w:lnNumType")
            ln.set(qn("w:countBy"), "1")
            ln.set(qn("w:restart"), "continuous")
            section._sectPr.append(ln)
        page_number(section)
    doc.save(path)


if __name__ == "__main__":
    main(sys.argv[1], line_numbers="--no-line-numbers" not in sys.argv)
