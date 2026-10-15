#!/usr/bin/env python3
"""canonical patient-variant CSV를 교수님 CRF 통합본 형식의 엑셀로 내보낸다.

- 시트: 유전자별 1개 + `_lists`
- 1행 제목, 2행 헤더(셀 코멘트에 입력 기준), 3행부터 patient-variant 1행
- Mutation Type: 내부 3분류(variant_type)를 CRF 2분류(Small mutation / Large CNV)로 합친다.
  variant_type이 없거나 unclear이면 CRF 헤더에 적힌 분류 기준을 적용한다.
- ClinVar ID: Variation ID를 쓰고 ClinVar variation 페이지 링크를 셀에 건다.
- Phenotype 라벨은 config/crf_phenotype_label_map.csv로 CRF 라벨에 맞춘다.
"""

import argparse
import csv
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "core"))
from timestamped_outputs import csv_timestamp, make_timestamp, resolve_csv_path  # noqa: E402

WORKSPACE_ROOT = Path(__file__).resolve().parents[2]
GUIDANCE_PATH = WORKSPACE_ROOT / "config" / "gene_phenotype_review_guidance.csv"
LABEL_MAP_PATH = WORKSPACE_ROOT / "config" / "crf_phenotype_label_map.csv"
MISSING = {"", "not specified", "n/a", "na", "none", "nan", "unknown"}
SUBCATEGORY_GENES = {"SYNGAP1"}

VARIANT_TYPE_TO_CRF = {
    "sequence_variant_or_small_indel": "Small mutation",
    "exon_level_cnv": "Large CNV",
    "large_cnv": "Large CNV",
}
# CRF 헤더 기준: exon deletion 이상, Mb/kb, microdel, arr, 좌표범위(genomic 5자리 이상)는 Large CNV
LARGE_CNV_PATTERN = re.compile(
    r"\bmb\b|\d\s*mb|\bkb\b|\d\s*kb|microdel|microdup|arr\[|\bex(on)?s?\b|\bex(on)?\s?\d|\bchr|g\.\(|"
    r"\d{5,}\s*[_\-–]\s*\d{5,}|\d{1,3}(,\d{3}){2,}|(?<![a-z])(\d{1,2}|x|y)[pq]\d|\bt\(|"
    r"\bdeletion\b|\bduplication\b|whole|entire|all gene|外显子|缺失|重复",
    re.IGNORECASE,
)
# point mutation / small indel로 볼 수 있는 표기 (g. 치환, IVS, 단백질식 표기, 짧은 del/dup/ins, frameshift 등)
SMALL_VARIANT_PATTERN = re.compile(
    r"[acgt]\s*>\s*[acgt]|[acgt]\(\d|ivs|int\d|^[a-z]\d+[a-z]$|del|dup|ins|fs|splice|"
    r"missens|nonsense|trunc|frameshift|mutation",
    re.IGNORECASE,
)


# CRF 통합본에서 치료 반응을 "효과 있는 치료 약어만" 적는 유전자 (나머지 유전자는 서술 그대로)
ABBREVIATED_RESPONSE = {
    "CDKL5": [("KD", r"ketogenic|\bKD\b|atkins"), ("CBD", r"cannabidiol|\bCBD\b"), ("FFL", r"fenfluramine|\bFFL\b")],
}
POSITIVE_RESPONSE = re.compile(r"benefi|effective|respon(se|der|sive)|improv|reduc|seizure[- ]free|control", re.I)
NEGATIVE_RESPONSE = re.compile(
    r"\bno\b|\bnot\b|without|ineffective|refractory|fail|unclear|cohort|aggregate|worse|aggravat|attempted|"
    r"non-?respon|unknown|discontinu", re.I)


def abbreviated_response(gene, value):
    """효과가 명시된 치료만 약어로 남긴다 (문장·세미콜론 단위, 부정 표현이 있으면 제외)."""
    found = []
    for clause in re.split(r"[;.]", "" if is_missing(value) else str(value)):
        for label, pattern in ABBREVIATED_RESPONSE[gene]:
            if (re.search(pattern, clause, re.I) and POSITIVE_RESPONSE.search(clause)
                    and not NEGATIVE_RESPONSE.search(clause) and label not in found):
                found.append(label)
    return ", ".join(found)


# CRF 통합본(0904) SCN1A 작성 방식: 표적 치료 반응이 비어 있으면 ASM 반응, 그것도 없으면 사용 약물을 적는다
FALLBACK_RESPONSE_GENES = {"SCN1A"}


def targeted_response_text(gene, r):
    value = r.get("targeted_treatment_response")
    if not is_missing(value) or gene not in FALLBACK_RESPONSE_GENES:
        return text(value)
    if not is_missing(r.get("response_to_asm")):
        return str(r.get("response_to_asm")).strip()
    if not is_missing(r.get("asm_used")):
        return f"drug(s) used: {str(r.get('asm_used')).strip()}"
    return "Not specified"


def is_missing(value):
    return pd.isna(value) or str(value).strip().lower() in MISSING


def text(value):
    return "Not specified" if is_missing(value) else str(value).strip()


def crf_mutation_type(row):
    """내부 variant_type을 CRF 2분류로 바꾸고, 없으면 CRF 헤더 기준으로 분류한다."""
    mapped = VARIANT_TYPE_TO_CRF.get(str(row.get("variant_type", "")).strip())
    if mapped:
        return mapped
    variant_c = "" if is_missing(row.get("variant_c")) else str(row.get("variant_c"))
    variant_p = "" if is_missing(row.get("variant_p")) else str(row.get("variant_p"))
    if re.search(r"wild[- ]type|no \w+ mutation|not available", variant_c, re.IGNORECASE):
        return ""
    if variant_p and variant_p != "p.?":
        return "Small mutation"
    if re.match(r"^(NM_[\d.]+:)?c\.", variant_c) and not re.search(r"\bex(on)?\s?\d", variant_c, re.IGNORECASE):
        return "Small mutation"
    if LARGE_CNV_PATTERN.search(variant_c):
        return "Large CNV"
    if SMALL_VARIANT_PATTERN.search(variant_c):
        return "Small mutation"
    return ""


def load_label_map():
    label_map = {}
    with LABEL_MAP_PATH.open(newline="", encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            label_map[(row["gene"], row["canonical_value"])] = (
                row["crf_phenotype"], row["crf_phenotype_subcategory"])
    return label_map


def load_guidance():
    with GUIDANCE_PATH.open(newline="", encoding="utf-8-sig") as f:
        return {row["gene"]: row for row in csv.DictReader(f)}


def header_comments(gene, guidance):
    g = guidance.get(gene, {})
    return {
        "Patient ID": "Unique patient identifier (DB field: canonical_patient_id). Same patient keeps the same ID across rows.",
        "Variant cDNA change": "Variant at cDNA level (DB field: variant_c). Non-standard genotype text is kept as written in the paper.",
        "Variant protein change": "Variant at protein level, HGVS 3-letter (DB field: variant_p)",
        "Mutation Type": "Small mutation = point mutation + small indel. Large CNV = exon-level deletion/duplication or larger CNV.",
        "ClinVar pathogenicity": "Pathogenic / Likely pathogenic / Pathogenic/Likely pathogenic only; blank otherwise (DB field: clinvar_pathogenicity)",
        "ClinVar ID": "ClinVar Variation ID; cell links to the ClinVar variation page (DB field: ClinVar ID)",
        "Source in PMID": "PMID of clinical information; cell links to PubMed (DB field: pmid)",
        "Sex": "Male / Female / Not specified",
        "Age at seizure onset": g.get("onset_unit", ""),
        "Age at last follow-up": g.get("followup_unit", ""),
        "Phenotype": "Select one: " + g.get("phenotype_categories", "").replace(" | ", " / "),
        "Phenotype_subcategory": "Subtype within Generalized/Other epilepsy",
        "VUS flag": "Variant classified as VUS (paper) or VUS/conflicting (ClinVar). Patient is kept; flag only.",
        "Response to targeted ASM": (
            "Only treatments with explicitly reported benefit, as abbreviations (KD / CBD / FFL); blank otherwise."
            if gene in ABBREVIATED_RESPONSE else g.get("response_focus", "").replace(" | ", "; ")),
    }


SUBCATEGORY_PATH = WORKSPACE_ROOT / "config" / "crf_phenotype_subcategory_values.csv"


def load_subcategories():
    """CRF에서 임상의가 넣은 하위분류 중 DB 표현형 분류에 대응이 없는 값 (gene, mention_id) -> 하위분류."""
    if not SUBCATEGORY_PATH.exists():
        return {}
    with SUBCATEGORY_PATH.open(newline="", encoding="utf-8-sig") as f:
        return {(row["gene"], row["mention_id"]): row["phenotype_subcategory"] for row in csv.DictReader(f)}


CLINVAR_DISPLAY_PATH = WORKSPACE_ROOT / "config" / "crf_clinvar_id_display.csv"


def load_clinvar_display():
    """CRF 통합본(0904)에서 임상의가 분류와 무관하게 적은 ClinVar ID (gene, mention_id) -> ID."""
    if not CLINVAR_DISPLAY_PATH.exists():
        return {}
    with CLINVAR_DISPLAY_PATH.open(newline="", encoding="utf-8-sig") as f:
        return {(row["gene"], row["mention_id"]): row["clinvar_id"] for row in csv.DictReader(f)}


def gene_rows(gene, df, label_map):
    rows = []
    subcategories = load_subcategories()
    clinvar_display = load_clinvar_display()
    for _, r in df.iterrows():
        category = text(r.get("phenotype_review_category"))
        phenotype, subcategory = label_map.get((gene, category), (category, ""))
        subcategory = subcategories.get((gene, str(r.get("mention_id", "")).split(".")[0]), subcategory)
        pmid = "" if is_missing(r.get("pmid")) else str(r.get("pmid")).split(".")[0]
        source = pmid
        if str(r.get("korean_patient_source", "")).strip() == "Korean multicenter cohort":
            source = f"{pmid}/ Korean cohort" if pmid else "Korean cohort"
        row = {
            "Patient ID": r.get("canonical_patient_id", ""),
            "Variant cDNA change": text(r.get("variant_c")),
            "Variant protein change": text(r.get("variant_p")),
            "Mutation Type": crf_mutation_type(r),
            "ClinVar pathogenicity": "" if is_missing(r.get("clinvar_pathogenicity")) else r.get("clinvar_pathogenicity"),
            "ClinVar ID": ("" if is_missing(r.get("ClinVar ID")) else str(r.get("ClinVar ID")).split(".")[0])
            or clinvar_display.get((gene, str(r.get("mention_id", "")).split(".")[0]), ""),
            "Source in PMID": source,
            "Sex": text(r.get("sex")).capitalize() if not is_missing(r.get("sex")) else "Not specified",
            "Age at seizure onset": text(r.get("age_of_onset")),
            "Age at last follow-up": text(r.get("last_followup_age")),
            "Phenotype": phenotype,
            "Phenotype_subcategory": subcategory,
            "Response to targeted ASM": (
                abbreviated_response(gene, r.get("targeted_treatment_response"))
                if gene in ABBREVIATED_RESPONSE else targeted_response_text(gene, r)),
            "_pmid": pmid,
        }
        if "variant_classification_flag" in df.columns:
            row["VUS flag"] = "" if is_missing(r.get("variant_classification_flag")) else r.get("variant_classification_flag")
        rows.append(row)
    return rows


def write_gene_sheet(wb, gene, rows, guidance, timestamp):
    columns = [
        "Patient ID", "Variant cDNA change", "Variant protein change", "Mutation Type",
        "ClinVar pathogenicity", "ClinVar ID", "Source in PMID", "Sex",
        "Age at seizure onset", "Age at last follow-up", "Phenotype",
    ]
    if gene in SUBCATEGORY_GENES:
        columns.append("Phenotype_subcategory")
    columns.append("Response to targeted ASM")
    if rows and "VUS flag" in rows[0]:
        columns.append("VUS flag")

    ws = wb.create_sheet(gene)
    ws.append([f"{gene} CRF  —  한 행에 환자-변이 1개. canonical {timestamp} 기준 자동 생성. "
               "같은 환자는 같은 Patient ID로 여러 행에 나올 수 있습니다."])
    ws.append(columns)
    comments = header_comments(gene, guidance)
    for cell in ws[2]:
        cell.font = Font(bold=True)
        if comments.get(cell.value):
            cell.comment = Comment(comments[cell.value], "pipeline")

    link_font = Font(color="0563C1", underline="single")
    for row in rows:
        ws.append([row[c] for c in columns])
        excel_row = ws.max_row
        clinvar_id = row["ClinVar ID"]
        if re.fullmatch(r"\d+", clinvar_id):
            cell = ws.cell(excel_row, columns.index("ClinVar ID") + 1)
            cell.hyperlink = f"https://www.ncbi.nlm.nih.gov/clinvar/variation/{clinvar_id}/"
            cell.font = link_font
        if row["_pmid"]:
            cell = ws.cell(excel_row, columns.index("Source in PMID") + 1)
            cell.hyperlink = f"https://pubmed.ncbi.nlm.nih.gov/{row['_pmid']}/"
            cell.font = link_font
    ws.freeze_panes = "B3"


def write_lists_sheet(wb, genes, guidance, label_map):
    lists = {"Sex": ["Male", "Female", "Not specified"]}
    for gene in genes:
        values = []
        for item in guidance.get(gene, {}).get("phenotype_categories", "").split(" | "):
            mapped = label_map.get((gene, item.strip()), (item.strip(), ""))[0]
            if mapped and mapped not in values:
                values.append(mapped)
        lists[f"{gene}_phen"] = values
    lists["MutType"] = ["Small mutation", "Large CNV"]
    ws = wb.create_sheet("_lists", 0)
    names = list(lists)
    ws.append(names)
    for i in range(max(len(v) for v in lists.values())):
        ws.append([lists[n][i] if i < len(lists[n]) else None for n in names])


def main():
    parser = argparse.ArgumentParser(description="canonical CSV -> CRF 형식 엑셀")
    parser.add_argument("--canonical-root", default="outputs/current/canonical_patients")
    parser.add_argument("--genes", nargs="+",
                        default=["SCN1A", "SCN2A", "KCNQ2", "CDKL5", "SYNGAP1", "EEF1A2",
                                 "PCDH19", "SCN8A", "GABRA1", "STXBP1", "SMC1A", "FOXG1", "KCNT1"])
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--timestamp", default=None)
    args = parser.parse_args()

    timestamp = make_timestamp(args.timestamp)
    canonical_root = WORKSPACE_ROOT / args.canonical_root
    guidance = load_guidance()
    label_map = load_label_map()

    wb = Workbook()
    wb.remove(wb.active)
    exported = []
    for gene in args.genes:
        path = resolve_csv_path(canonical_root / f"{gene}_canonical_patients", f"{gene}_canonical_patients")
        if path is None:
            print(f"{gene}: canonical CSV 없음 (건너뜀)")
            continue
        df = pd.read_csv(path, dtype=str)
        rows = gene_rows(gene, df, label_map)
        write_gene_sheet(wb, gene, rows, guidance, csv_timestamp(path))
        exported.append(gene)
        blank_type = sum(1 for r in rows if not r["Mutation Type"])
        print(f"{gene}: {len(rows)} rows from {path.name} (Mutation Type 미분류 {blank_type})")
    write_lists_sheet(wb, exported, guidance, label_map)

    output_dir = WORKSPACE_ROOT / args.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"Phenotype_CRF_{timestamp}.xlsx"
    wb.save(output_path)
    print(f"저장: {output_path}")


if __name__ == "__main__":
    main()
